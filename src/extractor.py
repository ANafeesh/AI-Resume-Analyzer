"""Document text extraction module supporting PDF, DOCX, and TXT files."""

from __future__ import annotations

import io
from pathlib import Path
from typing import BinaryIO, Union

import docx
import pdfplumber

from src.utils import clean_text, get_logger

logger = get_logger(__name__)

FileSource = Union[str, Path, BinaryIO, io.BytesIO]


def _get_bytes_stream(file_source: FileSource) -> tuple[io.BytesIO, str]:
    """Convert file source to BytesIO and determine the file suffix.

    Args:
        file_source: A path or file-like object.

    Returns:
        tuple[io.BytesIO, str]: A stream containing file bytes and the lowercase extension.
    """
    if isinstance(file_source, (str, Path)):
        path = Path(file_source)
        suffix = path.suffix.lower()
        with open(path, "rb") as f:
            data = f.read()
        return io.BytesIO(data), suffix

    # Streamlit UploadedFile or file-like object
    name = getattr(file_source, "name", "")
    suffix = Path(name).suffix.lower() if name else ""

    # Ensure pointer is at start
    if hasattr(file_source, "seek"):
        file_source.seek(0)
    data = file_source.read()
    if hasattr(file_source, "seek"):
        file_source.seek(0)

    return io.BytesIO(data), suffix


def extract_text_from_pdf(stream: io.BytesIO) -> tuple[str, str | None]:
    """Extract clean text from a PDF stream using pdfplumber.

    Args:
        stream: A binary stream containing PDF data.

    Returns:
        tuple[str, str | None]: Extracted text and error message if extraction fails.
    """
    try:
        stream.seek(0)
        with pdfplumber.open(stream) as pdf:
            if len(pdf.pages) == 0:
                return "", "The PDF file has no pages."

            extracted_pages: list[str] = []
            for i, page in enumerate(pdf.pages):
                try:
                    page_text = page.extract_text() or ""
                    extracted_pages.append(page_text)
                except Exception as page_err:
                    logger.warning("Failed to extract page %d: %s", i + 1, page_err)

            combined_text = "\n".join(extracted_pages).strip()
            if not combined_text:
                return (
                    "",
                    "The uploaded PDF appears to be a scanned document or image without a readable text layer. "
                    "Please upload a PDF with selectable text, a DOCX, or a TXT file.",
                )

            cleaned = clean_text(combined_text)
            return cleaned, None

    except Exception as e:
        logger.error("Corrupted or unreadable PDF: %s", e, exc_info=True)
        return "", f"Could not read PDF. The file may be damaged or password-protected ({e})."


def extract_text_from_docx(stream: io.BytesIO) -> tuple[str, str | None]:
    """Extract text from a DOCX stream including paragraphs and tables.

    Args:
        stream: A binary stream containing DOCX data.

    Returns:
        tuple[str, str | None]: Extracted text and error message if extraction fails.
    """
    try:
        stream.seek(0)
        doc = docx.Document(stream)
        text_parts: list[str] = []

        for p in doc.paragraphs:
            if p.text:
                text_parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append(" | ".join(row_cells))

        combined = "\n".join(text_parts).strip()
        if not combined:
            return "", "The DOCX document is empty."

        return clean_text(combined), None

    except Exception as e:
        logger.error("Corrupted or unreadable DOCX: %s", e, exc_info=True)
        return "", f"Could not read DOCX. The document may be corrupted or invalid ({e})."


def extract_text_from_txt(stream: io.BytesIO) -> tuple[str, str | None]:
    """Extract text from a TXT stream with encoding fallback.

    Args:
        stream: A binary stream containing text data.

    Returns:
        tuple[str, str | None]: Extracted text and error message if extraction fails.
    """
    try:
        stream.seek(0)
        raw_bytes = stream.read()
        if not raw_bytes:
            return "", "The text file is empty."

        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded = raw_bytes.decode(encoding)
                cleaned = clean_text(decoded)
                if not cleaned:
                    return "", "The text file contains only blank or unprintable characters."
                return cleaned, None
            except UnicodeDecodeError:
                continue

        return "", "Could not decode text file with standard encodings (UTF-8, Latin-1)."

    except Exception as e:
        logger.error("Error reading TXT file: %s", e, exc_info=True)
        return "", f"Error reading text file: {e}"


def extract_resume_text(
    file_source: FileSource,
    explicit_extension: str | None = None,
) -> tuple[str, str | None]:
    """Extract clean text from a PDF, DOCX, or TXT file source.

    Args:
        file_source: File path or file-like object.
        explicit_extension: Optional extension override (e.g., '.pdf').

    Returns:
        tuple[str, str | None]: (Extracted text, error message if any).
    """
    try:
        stream, detected_ext = _get_bytes_stream(file_source)
        extension = (explicit_extension or detected_ext or "").lower()

        # Check for empty file stream
        stream.seek(0, io.SEEK_END)
        size = stream.tell()
        stream.seek(0)
        if size == 0:
            return "", "Empty file uploaded (0 bytes). Please upload a valid resume."

        if extension == ".pdf":
            return extract_text_from_pdf(stream)
        elif extension == ".docx":
            return extract_text_from_docx(stream)
        elif extension == ".txt":
            return extract_text_from_txt(stream)
        else:
            msg = (
                f"Unsupported file type '{extension or 'unknown'}'. "
                "Supported formats are PDF (.pdf), Word (.docx), and Plain Text (.txt)."
            )
            logger.warning(msg)
            return "", msg

    except Exception as e:
        logger.error("Unexpected error extracting document: %s", e, exc_info=True)
        return "", f"An unexpected error occurred while processing the file: {e}"
